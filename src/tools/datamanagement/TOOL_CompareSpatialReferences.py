import os
import sys
import logging
import datetime
import pandas as pd
from pathlib import Path
from importlib import reload

import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

import arcpy
from arcgis.gis import GIS

sys.path.insert(0, str(Path(__file__).resolve().parents[3]))

from src.classes.datacatalog import DataCatalogRow
from src.constants.paths import PORTAL_ITEM_URL, LOG_DIR
from src.constants.values import SHEET_NAME
#################################################################################################################################################################################################################
## Globals
DATETIME_STR = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
#################################################################################################################################################################################################################
## Logging
reload(logging)
log_file =os.path.join(LOG_DIR, "CompareSpatialReferences",f"CompareSpatialReferences_{DATETIME_STR}.log")

logging.getLogger().disabled = True
logging.getLogger("arcgis.gis._impl._portalpy").setLevel(logging.WARNING)
logging.getLogger("urllib3.connectionpool").setLevel(logging.WARNING)
logging.getLogger("requests_oauthlib.oauth2_session").setLevel(logging.WARNING)

logger = logging.getLogger(__name__)
file_handler = logging.FileHandler(log_file, mode='w')
formatter = logging.Formatter('%(name)s - %(levelname)s - %(message)s')
file_handler.setFormatter(formatter)
logger.addHandler(file_handler)
logger.setLevel(logging.INFO)
#################################################################################################################################################################################################################

def excelFormatting(excel_path):
    logger.info(f"Formatting Excel Report...") 
    wb = openpyxl.load_workbook(excel_path)

    sheet_names = wb.sheetnames

    for sheet_name in sheet_names:
        ws = wb[sheet_name]

        if sheet_name in ws.tables:
            del ws.tables[sheet_name]
            
        table_dimensions = ws.calculate_dimension()
        
        tbl = Table(displayName=sheet_name, ref=table_dimensions)
        
        style = TableStyleInfo(name=f"TableStyleMedium2",showFirstColumn=True, showLastColumn=True, showRowStripes=True, showColumnStripes=False)
        
        tbl.tableStyleInfo = style
        ws.add_table(tbl)
        
        for l in ["A", "B", "C", "D", "E", "F", "G"]:
            if l == "A":
                ws.column_dimensions[l].width = 50
            elif l in ["B", "C"]:
                ws.column_dimensions[l].width = 25
            elif  l == "F":
                ws.column_dimensions[l].width = 90
            else:
                ws.column_dimensions[l].width = 175

    wb.save(excel_path)

    del wb


def getSpatialReferences(dataframe, gdb_path, gis_conn)->dict:
    output_dict = {}

    arcpy.AddMessage("Retrieving Spatial References...")

    for index, row in dataframe.iterrows():
        temp_dict ={}
        row_obj = DataCatalogRow(c_row=row, index=index, gdb_path=gdb_path, gis_conn=gis_conn)
        if row_obj.local_exist:
            file_path = row_obj.gdb_item_path
            if file_path:
                temp_dict["Local - Spatial Reference"] = arcpy.Describe(file_path).spatialReference.factoryCode
                temp_dict["Local - Path"] = file_path
            else:
                temp_dict["Local - Spatial Reference"] = None
                temp_dict["Local - Path"] = None

        if row_obj.service_exist:
            service_object = row_obj.getServiceObject()
            if service_object:
                temp_dict["Service - Spatial Reference"] = service_object.spatialReference
                temp_dict["Service - Portal URL"] = f"{PORTAL_ITEM_URL}{row_obj.agol_item_id}"
            else:
                temp_dict["Service - Spatial Reference"] = None
                temp_dict["Service - Portal URL"] = None
        
        output_dict[row_obj.table_name] = temp_dict

    return output_dict

def main(gis_conn:GIS, gdb_path:str, catalog_path:str, output_excel:str)->None:
    logger.info(f"Starting: {datetime.datetime.now()}")
    logger.info(f"Run by: {os.getlogin()}")
    logger.info(f"Run on: {datetime.datetime.now().strftime('%d/%m/%Y, %H:%M:%S')}")
    logger.info(f"GDB Directory: {gdb_path}")
    logger.info(f"Output Excel Path: {output_excel}")


    ## Converts the Data Catalog Excel to a Pandas DataFrame
    catalog_df = pd.read_excel(io=catalog_path, sheet_name=SHEET_NAME,header=0,names=None)

    df_dictionary = getSpatialReferences(catalog_df, gdb_path, gis_conn)
        
    arcpy.AddMessage("Building DataFrame...")
    df = pd.DataFrame.from_dict(df_dictionary, "index", columns=["Local - Spatial Reference", 
                                                                "Service - Spatial Reference",
                                                                "Local - Path",
                                                                "Service - Portal URL"])

    arcpy.AddMessage("Exporting Excel...")
    df.to_excel(output_excel, na_rep="N/A" ,index_label="Table Name")

    excelFormatting(excel_path=output_excel)

    arcpy.AddMessage(f"Excel Report Has Been Exported to:\n{output_excel}")
    arcpy.AddMessage(f"Opening Excel Report...")
    try:
        os.startfile(output_excel)
    except Exception as t:
        arcpy.AddWarning(f"Failed to Launch Excel")

    return


