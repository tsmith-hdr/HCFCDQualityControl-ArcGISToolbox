#################################################################################################################################################################################################################
## Libraries
import os
import sys
import itertools
import pandas as pd
import logging
from pathlib import Path
import datetime

import arcpy
from arcgis.gis import GIS

import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

if str(Path(__file__).resolve().parents[2]) not in sys.path:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))


from src.constants.paths import ROOT_DIR, PORTAL_ITEM_URL
from src.constants.values import SHEET_NAME
import src.classes.DataCatalog as dc
#################################################################################################################################################################################################################
## Logging ## Don't Change
#log_file = os.path.join(os.getcwd(),"ftp_achd.log")
log_file = Path(ROOT_DIR, "logs", "CompareStorageLocations","CompareStorageLocations.log")
logging.basicConfig(filename=log_file, filemode="w",format='%(name)s - %(levelname)s - %(message)s', level=logging.INFO)
logging.getLogger("arcgis.gis._impl._portalpy").setLevel(logging.WARNING)
logging.getLogger("urllib3.connectionpool").setLevel(logging.WARNING)
logging.getLogger("requests_oauthlib.oauth2_session").setLevel(logging.WARNING)
logger = logging.getLogger(__name__)
#################################################################################################################################################################################################################
#################################################################################################################################################################################################################
## Functions
def createServicesList(gis_conn:GIS)->list:
    """
    This function creates a list of all Feature Layer and Map Image Layer services in the AGOL Portal
    Args: GIS Connection
    Returns: List of AGOL Services
    """
    service_item_types = ["Feature Layer", "Map Image Layer"]
    logger.info(f"Creating Services List...")
    services_list = []
    for service_item_type in service_item_types:
        result_list = gis_conn.content.search(query="", item_type=service_item_type, max_items=-1)

        services_list = services_list + result_list

    return services_list


def createLocalList(gdb_path:Path)->list:
    """
    This function creates a list of all the Feature Classes and Rasters in the Master File GDB
    Args: Master File GDB
    Returns: List of Feature Classes and Rasters in the Master File GDB
    """
    logger.info(f"Creating Local File GDB List...")

    arcpy.env.workspace = str(gdb_path)

    if arcpy.env.workspace != str(gdb_path):
        logger.error("!! Workspaces Don't Match !!")
        sys.exit("!! Workspaces Don't Match !!")
    gdb_item_list = []
    datasets = arcpy.ListDatasets(feature_type="Feature")
    datasets.append(None)
    [gdb_item_list.append(featureclass) for dataset in datasets for  featureclass in arcpy.ListFeatureClasses(feature_dataset=dataset)]
    [gdb_item_list.append(raster) for raster in arcpy.ListRasters()]


    logger.debug(gdb_item_list)

    return gdb_item_list


def excelFormatting(excel_path):
    logger.info(f"Formatting Excel Report...") 
    wb = openpyxl.load_workbook(excel_path)

    sheet_names = wb.sheetnames

    for sheet_name in sheet_names:
        ws = wb[sheet_name]

        if sheet_name in ws.tables:
            del ws.tables[sheet_name]
            
        table_dimensions = ws.calculate_dimension()
        
        tbl = Table(displayName=sheet_name, ref=table_dimensions)
        
        style = TableStyleInfo(name=f"TableStyleMedium2",showFirstColumn=True, showLastColumn=True, showRowStripes=True, showColumnStripes=False)
        
        tbl.tableStyleInfo = style
        ws.add_table(tbl)
        
        for l in ["A", "B", "C", "D", "E", "F", "G"]:
            if l == "A":
                ws.column_dimensions[l].width = 50
            elif l in ["B", "C", "D"]:
                ws.column_dimensions[l].width = 35
            elif  l == "G":
                ws.column_dimensions[l].width = 90
            else:
                ws.column_dimensions[l].width = 175

    wb.save(excel_path)

    del wb



#################################################################################################################################################################################################################
def main(gis_conn:GIS, gdb_path:str, catalog_path:str, excel_path:str)->None:
    logger.info(f"Starting: {datetime.datetime.now()}")
    logger.info(f"Run by: {os.getlogin()}")
    logger.info(f"Run on: {datetime.datetime.now().strftime('%d/%m/%Y, %H:%M:%S')}")
    logger.info(f"Excel Report: {excel_path}")
    logger.info(f"Data Catalog: {catalog_path}")
    logger.info(f"FGDB Path: {gdb_path}")

    df_list = []

    catalog_df = pd.read_excel(io=catalog_path, sheet_name=SHEET_NAME,header=0,names=None)

    local_list = createLocalList(gdb_path=gdb_path)

    service_list = createServicesList(gis_conn=gis_conn)
    service_id_list = [i.id for i in service_list]


    for index, row in catalog_df.iterrows():
        temp_dict = {}
        row_obj = dc.DataCatalogRow(c_row=row, index=index, gdb_path=gdb_path, gis_conn=gis_conn)
        
        temp_dict["Table Name"] = row_obj.table_name
        temp_dict["Data Catalog - Exist"] = True
        temp_dict["Local - Exist"] = row_obj.local_exist
        temp_dict["Service - Exist"] = row_obj.service_exist

        if row_obj.local_exist:
            local_list.remove(row_obj.table_name)
            temp_dict["Local - Path"] = row_obj.gdb_item_path
        if row_obj.service_exist:
            service_id_list.remove(row_obj.agol_item_id)
            temp_dict["Service - Item Id"] = row_obj.agol_item_id
            temp_dict["Service - Path"] = f"{PORTAL_ITEM_URL}{row_obj.agol_item_id}"
        
        df_list.append(temp_dict)

    for table_name in local_list:
        temp_dict = {}
        temp_dict["Table Name"] = table_name
        temp_dict["Data Catalog - Exist"] = False
        temp_dict["Local - Exist"] = True
        temp_dict["Service - Exist"] = "No Associated Service in Data Catalog"
        temp_dict["Local - Path"] = os.path.join(gdb_path, table_name)
    
        df_list.append(temp_dict)

    for item_id in service_id_list:
        temp_dict = {}
        temp_dict["Service - Item Id"] = item_id
        temp_dict["Service - Path"] = f"{PORTAL_ITEM_URL}{item_id}"
        temp_dict["Data Catalog - Exist"] = False
        temp_dict["Local - Exist"] = "No Associated GDB Item in Data Catalog"
        temp_dict["Service - Exist"] = True

        df_list.append(temp_dict)
    

    df = pd.DataFrame(df_list, columns=["Table Name",
                                        "Data Catalog - Exist",
                                        "Local - Exist",
                                        "Service - Exist",
                                        "Local - Path",
                                        "Service - Path",
                                        "Service - Item Id"]
                                        )


    df.to_excel(excel_path, sheet_name="StorageComparison", index=False)

    excelFormatting(excel_path)

    arcpy.AddMessage(f"Excel Report Has Been Exported to:\n{excel_path}")
    arcpy.AddMessage(f"Opening Excel Report...")
    try:
        os.startfile(excel_path)
    except Exception as t:
        arcpy.AddWarning(f"Failed to Launch Excel")


    return