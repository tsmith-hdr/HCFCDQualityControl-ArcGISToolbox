from arcgis.gis import GIS
import arcpy
import pandas as pd
from pathlib import Path
import logging
import sys
import datetime
import os
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

if str(Path(__file__).resolve().parents[1]) not in sys.path:
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from hcfcd_classes.DataCatalog import DataCatalogRow
from hcfcd_constants.values import ROOT_DIR

#################################################################################################################################################################################################################
## Logging ## Don't Change
log_file = Path(ROOT_DIR, "logs", "CompareSpatialReferences", "CompareSpatialReferences.log")
logging.basicConfig(filename=log_file, filemode="w",format='%(name)s - %(levelname)s - %(message)s', level=logging.INFO)
logging.getLogger("arcgis.gis._impl._portalpy").setLevel(logging.WARNING)
logging.getLogger("urllib3.connectionpool").setLevel(logging.WARNING)
logging.getLogger("requests_oauthlib.oauth2_session").setLevel(logging.WARNING)
logger = logging.getLogger(__name__)
#################################################################################################################################################################################################################


def main(gis_conn:GIS, gdb_directory:Path, catalog_path:Path, output_excel:Path)->None:
    logger.info(f"Starting: {datetime.datetime.now()}")
    logger.info(f"Run by: {os.getlogin()}")
    logger.info(f"Run on: {datetime.datetime.now().strftime('%d/%m/%Y, %H:%M:%S')}")
    logger.info(f"GDB Directory: {gdb_directory}")
    logger.info(f"Output Excel Path: {output_excel}")

    
    DF_DICTIONARY = {}

    ## Converts the Data Catalog Excel to a Pandas DataFrame
    catalog_df = pd.read_excel(io=catalog_path, sheet_name="Inventory",header=0,names=None)
    arcpy.AddMessage("Retrieving Spatial References...")
    for index, row in catalog_df.iterrows():
        out_dict ={}
        row_obj = DataCatalogRow(row, index, gdb_directory, "Master", gis_conn)
        if row_obj.master_exist:
            file_path = row_obj.getFilePath("Master")
            if file_path:
                out_dict["Master - Spatial Reference"] = arcpy.Describe(file_path).spatialReference.factoryCode
                out_dict["Master - Path"] = file_path
            else:
                out_dict["Master - Spatial Reference"] = None
                out_dict["Master - Path"] = None
        if row_obj.spatial_exist:
            file_path = row_obj.getFilePath("Spatial")
            if file_path:
                out_dict["Spatial - Spatial Reference"] = arcpy.Describe(file_path).spatialReference.factoryCode
                out_dict["Spatial - Path"] = file_path
            else:
                out_dict["Spatial - Spatial Reference"] = None
                out_dict["Spatial - Path"] = None
                
        if row_obj.service_exist:
            service_object = row_obj.getServiceObject()
            if service_object:
                out_dict["Service - Spatial Reference"] = service_object.spatialReference
                out_dict["Service - Portal URL"] = row_obj.agol_link
            else:
                out_dict["Service - Spatial Reference"] = None
                out_dict["Service - Portal URL"] = None
        
        DF_DICTIONARY[row_obj.table_name] = out_dict
        
    arcpy.AddMessage("Building DataFrame...")
    df = pd.DataFrame.from_dict(DF_DICTIONARY, "index", columns=["Master - Spatial Reference", 
                                                                "Spatial - Spatial Reference", 
                                                                "Service - Spatial Reference",
                                                                "Master - Path",
                                                                "Spatial - Path",
                                                                "Service - Portal URL"])

    arcpy.AddMessage("Exporting Excel...")
    df.to_excel(output_excel, na_rep="N/A" ,index_label="Table Name")


    logger.info(f"Formatting Excel Report...") 
    wb = openpyxl.load_workbook(output_excel)

    sheet_names = wb.sheetnames

    for sheet_name in sheet_names:
        ws = wb[sheet_name]

        if sheet_name in ws.tables:
            del ws.tables[sheet_name]
            
        table_dimensions = ws.calculate_dimension()
        
        tbl = Table(displayName=sheet_name, ref=table_dimensions)
        
        style = TableStyleInfo(name=f"TableStyleMedium2",showFirstColumn=True, showLastColumn=True, showRowStripes=True, showColumnStripes=False)
        
        tbl.tableStyleInfo = style
        ws.add_table(tbl)
        
        for l in ["A", "B", "C", "D", "E", "F", "G"]:
            if l == "A":
                ws.column_dimensions[l].width = 50
            elif l in ["B", "C", "D"]:
                ws.column_dimensions[l].width = 25
            elif  l == "G":
                ws.column_dimensions[l].width = 90
            else:
                ws.column_dimensions[l].width = 175

    wb.save(output_excel)

    del wb

    return


