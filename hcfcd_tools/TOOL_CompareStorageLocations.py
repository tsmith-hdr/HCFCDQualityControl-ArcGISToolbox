#################################################################################################################################################################################################################
## Libraries
import os
import sys
import itertools
import pandas as pd
import logging
from pathlib import Path
import datetime

import arcpy
from arcgis.gis import GIS

import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.append(Path(__file__).resolve().parents[1])

from hcfcd_constants.values import ROOT_DIR, PORTAL_ITEM_URL
#################################################################################################################################################################################################################
## Logging ## Don't Change
#log_file = os.path.join(os.getcwd(),"ftp_achd.log")
log_file = Path(ROOT_DIR, "logs", "CompareStorageLocations","CompareStorageLocations.log")
logging.basicConfig(filename=log_file, filemode="w",format='%(name)s - %(levelname)s - %(message)s', level=logging.INFO)
logging.getLogger("arcgis.gis._impl._portalpy").setLevel(logging.WARNING)
logging.getLogger("urllib3.connectionpool").setLevel(logging.WARNING)
logging.getLogger("requests_oauthlib.oauth2_session").setLevel(logging.WARNING)
logger = logging.getLogger(__name__)
#################################################################################################################################################################################################################
#################################################################################################################################################################################################################
## Functions
def createServicesList(gis_conn:GIS)->list:
    """
    This function creates a list of all Feature Layer and Map Image Layer services in the AGOL Portal
    Args: GIS Connection
    Returns: List of AGOL Services
    """
    service_item_types = ["Feature Layer", "Map Image Layer"]
    logger.info(f"Creating Services List...")
    services_list = []
    for service_item_type in service_item_types:
        result_list = gis_conn.content.search(query="", item_type=service_item_type, max_items=-1)

        services_list = services_list + result_list

    return services_list

def createSpatialDictionary(gdb_dir:Path, avoidance_list:list)->dict:
    """
    Function creates a dictionary of all the Feature Classes and Rasters in the Spatial File GDBs
    Args: Directory of Spatial File GDBs, List of GDBs to Avoid
    Returns: Dictionary of GDBs and their respective Feature Classes and Rasters
    """
    logger.info(f"Creating Spatial File GDB List...")

    spatial_gdb_dictionary = {}

    spatial_gdb_list = [os.path.join(gdb_dir, g) for g in os.listdir(gdb_dir) if g.endswith('.gdb') and g not in avoidance_list]

    for gdb in spatial_gdb_list:
        spatial_gdb_dictionary[gdb] = []

        arcpy.env.workspace = gdb
        if arcpy.env.workspace != gdb:
            logger.error("!! Workspaces Don't Match !!")
            sys.exit("!! Workspaces Don't Match !!")

        fc_list = arcpy.ListFeatureClasses()
        raster_list = arcpy.ListRasters()
        item_list = fc_list + raster_list
        logger.debug(item_list)

        spatial_gdb_dictionary[gdb] = item_list

    return spatial_gdb_dictionary

def createMasterList(master_gdb:Path)->list:
    """
    This function creates a list of all the Feature Classes and Rasters in the Master File GDB
    Args: Master File GDB
    Returns: List of Feature Classes and Rasters in the Master File GDB
    """
    logger.info(f"Creating Master File GDB List...")

    arcpy.env.workspace = str(master_gdb)

    if arcpy.env.workspace != str(master_gdb):
        logger.error("!! Workspaces Don't Match !!")
        sys.exit("!! Workspaces Don't Match !!")

    fc_list = arcpy.ListFeatureClasses()
    raster_list = arcpy.ListRasters()

    master_list = fc_list + raster_list
    logger.debug(master_list)

    return master_list



def createCombinedDictionary(master_list:list, spatial_dictionary:dict, master_gdb:str)->dict:
    """
    This function combines the Master List and the Spatial Dictionary into a single dictionary
    Args: Master List, Spatial Dictionary, Master GDB Name
    Returns: Combined Dictionary
    """
    updated_spatial_dictionary = {}
    for k,v in spatial_dictionary.items():
        updated_spatial_dictionary[k] = v
    updated_spatial_dictionary[master_gdb] = master_list

    return updated_spatial_dictionary


def checkItem(item:str, item_list:list, item_type:str, check_list:list)->str:
    if item in check_list and item in item_list:
        return "Both"
    
    elif item in check_list and item not in item_list:
        if item_type != "MastervSpatial":
            return "Data Catalog Only"
        else:
            return "Master GDB Only"
        
    elif item not in check_list and item in item_list:
        if item_type != "MastervSpatial":
            return "Data Source Only"
        else:
            return "Spatial GDB Only"
        
    else:
        return False


#################################################################################################################################################################################################################
def main(gis_conn:GIS, gdb_directory:Path, catalog_path:Path, excel_path:Path)->None:
    logger.info(f"Starting: {datetime.datetime.now()}")
    logger.info(f"Run by: {os.getlogin()}")
    logger.info(f"Run on: {datetime.datetime.now().strftime('%d/%m/%Y, %H:%M:%S')}")
    logger.info(f"Excel Report: {excel_path}")
    logger.info(f"Data Catalog: {catalog_path}")
    logger.info(f"FGDB Directory: {gdb_directory}")

    catalog_df = pd.read_excel(io=catalog_path, sheet_name="Inventory",header=0,names=None)


    services_df = catalog_df[catalog_df['AGOL Link'].str.contains("http", na=False)]
    services_df["Item_Tuple"] = list(zip(services_df["AGOL Link"], services_df["GDB File Name"]))

    logger.info(f"Building Input Dictionary...")

    services_list = createServicesList(gis_conn)

    spatial_dictionary = createSpatialDictionary(gdb_directory, avoidance_list=["Master.gdb"])
    spatial_list = list(set(itertools.chain(*list(spatial_dictionary.values()))))

    master_gdb = Path(gdb_directory, "Master.gdb")
    master_list = createMasterList(master_gdb)

    combined_dictionary = createCombinedDictionary(master_list, spatial_dictionary, master_gdb=master_gdb)
    combined_list = list(set(spatial_list + master_list))


    input_dict = {"Service":{"item_list":services_list, "dataframe":services_df, "df_column":"Item_Tuple"}, 
                "Spatial":{"item_list":spatial_dictionary, "dataframe":catalog_df, "df_column":"GDB File Name"}, 
                "Master":{"item_list":master_list, "dataframe":catalog_df, "df_column":"GDB File Name"},
                "MastervSpatial":{"item_list":combined_dictionary, "dataframe":catalog_df, "df_column":"GDB File Name"}
                }



    out_dfs = []
    

    for item_type in list(input_dict.keys()):
        logger.info(f"---- {item_type} ----")

        df_list = []

        item_list = input_dict[item_type]["item_list"]
        pandas_list = input_dict[item_type]["dataframe"][input_dict[item_type]["df_column"]].tolist()

        if item_type == "Service":
            _checked = []
            pandas_id_list = []
            for p in pandas_list:
                agol_link = p[0]
                item_id = agol_link.split("=")[1]
                cleaned_item_id = item_id.replace("#overview","")
                pandas_id_list.append(cleaned_item_id)
            #pandas_id_list = [p[0].split("=")[1].replace("#overview","")[0] for p in pandas_list] ## Retrieves just the Item Id from the AGOL Portal Item URL
            id_list = [i.id for i in item_list] ## Creates a list of Item Ids from the items on the AGOL Portal
            print(id_list)
            for service in item_list:
                df_list.append([service.title, f"{PORTAL_ITEM_URL}{service.id}", checkItem(service.id, id_list, item_type, pandas_id_list)])
                _checked.append(f"{PORTAL_ITEM_URL}{service.id}")
            print(_checked)
            for url, name in pandas_list:
                if url.replace("#overview", "") not in _checked:
                    df_list.append([name, url, "Data Catalog Only"])



        elif item_type == "Spatial":
            _checked = []
            for spatial_gdb, table_name_list in item_list.items():
                logger.info(f"--{spatial_gdb}")
                for table_name in table_name_list:

                    df_list.append([table_name, os.path.basename(spatial_gdb), checkItem(table_name, table_name_list, item_type, pandas_list)])
                    _checked.append(table_name)
            
            for table_name in master_list:
                if table_name not in _checked:
                    df_list.append([table_name, os.path.basename(master_gdb), checkItem(table_name, master_list, item_type, pandas_list)])
                    _checked.append(table_name)

            for table_name in pandas_list:
                if table_name not in _checked:
                    df_list.append([table_name, "Not In GDB", checkItem(table_name, combined_list, item_type, pandas_list)])
        
        elif item_type == "Master":
            combined_list = list(set(itertools.chain(pandas_list, item_list)))
            for table_name in combined_list:

                df_list.append([table_name, os.path.basename(master_gdb), checkItem(table_name, item_list, item_type, pandas_list)]) 

        elif item_type == "MastervSpatial":
            _checked = []
            spatial_list = list(set(itertools.chain(*list(spatial_dictionary.values()))))
            master_list = input_dict["Master"]["item_list"]
            for gdb, table_name_list in spatial_dictionary.items():
                logger.info(f"--{gdb}")

                for table_name in table_name_list:

                    df_list.append([table_name, os.path.basename(gdb), checkItem(table_name, spatial_list, item_type, master_list)])
                    _checked.append(table_name)

            for table_name in master_list:
                if table_name not in _checked:
                    df_list.append([table_name, os.path.basename(master_gdb), checkItem(table_name, spatial_list, item_type, master_list)])
                    _checked.append(table_name)

            for table_name in pandas_list:
                if table_name not in _checked:
                    df_list.append([table_name, "Not In GDB", checkItem(table_name, combined_list, item_type, pandas_list)])



        df = pd.DataFrame(df_list, columns=["Item Name", "Source","Existance"])
        
        out_dfs.append((df, item_type))


    logger.info(f"Exporting DataFrames to Excel...")
    with pd.ExcelWriter(excel_path) as writer:
        out_dfs[0][0].to_excel(writer, sheet_name=out_dfs[0][1], index=False)
        out_dfs[1][0].to_excel(writer, sheet_name=out_dfs[1][1], index=False)
        #out_dfs[2][0].to_excel(writer, sheet_name=out_dfs[2][1], index=False)
        out_dfs[3][0].to_excel(writer, sheet_name=out_dfs[3][1], index=False)

        
    logger.info(f"Formatting Excel Report...") 
    wb = openpyxl.load_workbook(excel_path)

    sheet_names = wb.sheetnames
    count = 1
    for sheet_name in sheet_names:
        ws = wb[sheet_name]

        
        if sheet_name in ws.tables:
            del ws.tables[sheet_name]
            
        table_dimensions = ws.calculate_dimension()
        
        tbl = Table(displayName=sheet_name, ref=table_dimensions)
        
        style = TableStyleInfo(name=f"TableStyleMedium{count+1}",showFirstColumn=True, showLastColumn=True, showRowStripes=True, showColumnStripes=False)
        
        tbl.tableStyleInfo = style
        ws.add_table(tbl)
        
        ws.column_dimensions["A"].width = 80
        ws.column_dimensions["B"].width = 100
        ws.column_dimensions["C"].width = 20
        

        count+=1
    wb.save(excel_path)

    del wb

    return